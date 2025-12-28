"""
AI 트렌드 2025 PPT 생성 스크립트
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as RgbColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def add_title_slide(prs, title, subtitle):
    """타이틀 슬라이드 추가"""
    slide_layout = prs.slide_layouts[6]  # 빈 슬라이드
    slide = prs.slides.add_slide(slide_layout)

    # 배경색 설정
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RgbColor(30, 58, 138)  # 진한 파랑
    background.line.fill.background()

    # 제목
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(54)
    title_para.font.bold = True
    title_para.font.color.rgb = RgbColor(255, 255, 255)
    title_para.alignment = PP_ALIGN.CENTER

    # 부제목
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.text = subtitle
    subtitle_para.font.size = Pt(28)
    subtitle_para.font.color.rgb = RgbColor(200, 200, 200)
    subtitle_para.alignment = PP_ALIGN.CENTER

    return slide

def add_content_slide(prs, title, content_items, accent_color=RgbColor(59, 130, 246)):
    """내용 슬라이드 추가"""
    slide_layout = prs.slide_layouts[6]  # 빈 슬라이드
    slide = prs.slides.add_slide(slide_layout)

    # 상단 바
    top_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2)
    )
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = accent_color
    top_bar.line.fill.background()

    # 제목
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = title
    title_para.font.size = Pt(36)
    title_para.font.bold = True
    title_para.font.color.rgb = RgbColor(255, 255, 255)

    # 내용
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(8.6), Inches(5))
    content_frame = content_box.text_frame
    content_frame.word_wrap = True

    for i, item in enumerate(content_items):
        if i == 0:
            para = content_frame.paragraphs[0]
        else:
            para = content_frame.add_paragraph()

        para.text = f"• {item}"
        para.font.size = Pt(22)
        para.font.color.rgb = RgbColor(50, 50, 50)
        para.space_after = Pt(16)
        para.level = 0

    return slide

def add_section_slide(prs, section_title, section_number):
    """섹션 구분 슬라이드"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    # 배경
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RgbColor(16, 185, 129)  # 에메랄드
    background.line.fill.background()

    # 섹션 번호
    num_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1))
    num_frame = num_box.text_frame
    num_para = num_frame.paragraphs[0]
    num_para.text = f"0{section_number}"
    num_para.font.size = Pt(72)
    num_para.font.bold = True
    num_para.font.color.rgb = RgbColor(255, 255, 255)
    num_para.alignment = PP_ALIGN.CENTER

    # 섹션 제목
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(9), Inches(1))
    title_frame = title_box.text_frame
    title_para = title_frame.paragraphs[0]
    title_para.text = section_title
    title_para.font.size = Pt(40)
    title_para.font.color.rgb = RgbColor(255, 255, 255)
    title_para.alignment = PP_ALIGN.CENTER

    return slide

def create_ai_trends_ppt():
    """AI 트렌드 2025 PPT 생성"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # 1. 타이틀 슬라이드
    add_title_slide(prs, "AI 트렌드 2025", "인공지능이 바꾸는 미래")

    # 2. 목차
    add_content_slide(prs, "📋 목차", [
        "생성형 AI의 진화",
        "멀티모달 AI",
        "AI 에이전트",
        "엣지 AI & 온디바이스 AI",
        "AI 규제와 윤리",
        "산업별 AI 적용 사례",
        "미래 전망"
    ], RgbColor(107, 114, 128))

    # 3. 생성형 AI 섹션
    add_section_slide(prs, "생성형 AI의 진화", 1)

    add_content_slide(prs, "🤖 생성형 AI (Generative AI)", [
        "GPT-4, Claude, Gemini 등 대형 언어모델의 지속적 발전",
        "텍스트, 이미지, 비디오, 음악 등 다양한 콘텐츠 생성",
        "코드 생성 및 소프트웨어 개발 자동화",
        "기업용 맞춤형 AI 솔루션 확대",
        "RAG (Retrieval-Augmented Generation) 기술 고도화"
    ], RgbColor(139, 92, 246))

    # 4. 멀티모달 AI 섹션
    add_section_slide(prs, "멀티모달 AI", 2)

    add_content_slide(prs, "🎯 멀티모달 AI", [
        "텍스트 + 이미지 + 음성 + 비디오 통합 처리",
        "더 자연스러운 인간-AI 상호작용",
        "실시간 번역 및 통역 서비스",
        "의료 영상 분석과 진단 보조",
        "자율주행 및 로보틱스 분야 활용"
    ], RgbColor(236, 72, 153))

    # 5. AI 에이전트 섹션
    add_section_slide(prs, "AI 에이전트", 3)

    add_content_slide(prs, "🚀 AI 에이전트 (Agentic AI)", [
        "자율적으로 작업을 계획하고 실행하는 AI",
        "복잡한 업무 프로세스 자동화",
        "웹 브라우징, 파일 관리, 코딩 등 수행",
        "인간과 AI의 협업 워크플로우",
        "2025년 가장 주목받는 AI 트렌드"
    ], RgbColor(245, 158, 11))

    # 6. 엣지 AI 섹션
    add_section_slide(prs, "엣지 AI & 온디바이스 AI", 4)

    add_content_slide(prs, "📱 엣지 AI & 온디바이스 AI", [
        "스마트폰, IoT 기기에서 직접 AI 실행",
        "클라우드 의존도 감소, 프라이버시 강화",
        "실시간 처리 및 낮은 지연시간",
        "Apple Intelligence, Google Gemini Nano 등",
        "소형화된 언어모델 (SLM) 발전"
    ], RgbColor(6, 182, 212))

    # 7. AI 규제와 윤리 섹션
    add_section_slide(prs, "AI 규제와 윤리", 5)

    add_content_slide(prs, "⚖️ AI 규제와 윤리", [
        "EU AI Act 시행 (2024년~)",
        "AI 투명성 및 설명가능성 요구 증가",
        "딥페이크 및 허위정보 대응",
        "AI 저작권 및 지적재산권 논의",
        "책임있는 AI (Responsible AI) 프레임워크"
    ], RgbColor(239, 68, 68))

    # 8. 산업별 적용 사례
    add_section_slide(prs, "산업별 AI 적용", 6)

    add_content_slide(prs, "🏭 산업별 AI 적용 사례", [
        "헬스케어: 신약 개발, 진단 보조, 맞춤형 치료",
        "금융: 사기 탐지, 리스크 관리, 로보어드바이저",
        "제조: 예측 정비, 품질 관리, 스마트 팩토리",
        "교육: 개인화 학습, AI 튜터, 평가 자동화",
        "리테일: 추천 시스템, 재고 관리, 고객 서비스"
    ], RgbColor(34, 197, 94))

    # 9. 미래 전망
    add_section_slide(prs, "미래 전망", 7)

    add_content_slide(prs, "🔮 AI의 미래 전망", [
        "AGI (범용 인공지능)를 향한 지속적 발전",
        "AI와 인간의 공존 및 협력 모델 정립",
        "새로운 일자리 창출과 기존 직업의 변화",
        "AI 민주화: 누구나 AI를 활용하는 시대",
        "지속가능한 AI: 에너지 효율성 개선"
    ], RgbColor(99, 102, 241))

    # 10. 마무리 슬라이드
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RgbColor(30, 58, 138)
    background.line.fill.background()

    thanks_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.8), Inches(9), Inches(1.5))
    thanks_frame = thanks_box.text_frame
    thanks_para = thanks_frame.paragraphs[0]
    thanks_para.text = "감사합니다"
    thanks_para.font.size = Pt(54)
    thanks_para.font.bold = True
    thanks_para.font.color.rgb = RgbColor(255, 255, 255)
    thanks_para.alignment = PP_ALIGN.CENTER

    contact_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(9), Inches(0.8))
    contact_frame = contact_box.text_frame
    contact_para = contact_frame.paragraphs[0]
    contact_para.text = "AI 트렌드 2025"
    contact_para.font.size = Pt(24)
    contact_para.font.color.rgb = RgbColor(200, 200, 200)
    contact_para.alignment = PP_ALIGN.CENTER

    # 파일 저장
    output_path = "AI_Trends_2025.pptx"
    prs.save(output_path)
    print(f"PPT 파일이 생성되었습니다: {output_path}")
    return output_path

if __name__ == "__main__":
    create_ai_trends_ppt()
